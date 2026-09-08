import os
from datetime import datetime
import pyautogui
import pytesseract
import keyboard
from openpyxl import Workbook, load_workbook

# Point to Tesseract binary path
pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

# Define capture region: (left, top, width, height)
CAPTURE_ZONE = (1290, 973, 48, 19)

wavelength = input("Enter Wavelength(nm) to check Intensity: ")
start_val = float(input("Enter starting value: "))
step_val = float(input("Enter step value: "))

EXCEL_FILE = f"{wavelength}.xlsx"
current_step = start_val


def save_to_excel(data):
    global current_step
    try:
        # Create Excel file if it does not exist
        if not os.path.exists(EXCEL_FILE):
            wb = Workbook()
            ws = wb.active
            ws.append(["Timestamp", "Step Value", "Reading"])
            wb.save(EXCEL_FILE)
            print(f"File created: {os.path.abspath(EXCEL_FILE)}")

        # Append row to existing file
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active

        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.append([timestamp, round(current_step, 2), data])
        wb.save(EXCEL_FILE)

        print(f"Captured: {data} at Step: {round(current_step, 2)}")

        # Increment step value for next hotkey trigger
        current_step += step_val

    except Exception as e:
        print(f"Error saving to Excel: {e}")


def capture_logic():
    try:
        print("Capturing...")
        screenshot = pyautogui.screenshot(region=CAPTURE_ZONE)
        screenshot = screenshot.convert('L')  # Convert to grayscale for improved OCR

        extracted_text = pytesseract.image_to_string(screenshot, config='--psm 6').strip()

        if extracted_text:
            save_to_excel(extracted_text)
        else:
            print("Detected blank text. Check if CAPTURE_ZONE coordinates are correct.")

    except Exception as e:
        print(f"\nOCR Error: {e}")


if __name__ == "__main__":
    print(f"\nScript Active. Saving to {EXCEL_FILE}.")
    print("Script is running. Press F9 to capture. Press 'Esc' to stop.")
    
    keyboard.add_hotkey('F9', capture_logic)
    keyboard.wait('esc')
